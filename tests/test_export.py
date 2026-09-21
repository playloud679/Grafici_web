"""Tests for raw XLS export frame building."""
import sys
import unittest
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import export


def _df(times, values):
    return pd.DataFrame({"Time": pd.to_datetime(times), "value": values})


class TestBuildExportFrame(unittest.TestCase):
    def test_sample_numbers_in_time_order(self):
        ds = {
            "data": _df(["2026-01-02 10:00", "2026-01-01 10:00"], [2, 1]),
            "enabled": True,
        }
        frame = export.build_export_frame([ds], {"trim_start": 0, "trim_end": 0})
        self.assertEqual(list(frame["value"]), [1, 2])
        self.assertEqual(list(frame["Campione"]), [1, 2])

    def test_trim_removes_initial_and_final_rows(self):
        times = pd.date_range("2026-01-01", periods=10, freq="15min")
        ds = {"data": _df(times, list(range(10))), "enabled": True}
        frame = export.build_export_frame([ds], {"trim_start": 2, "trim_end": 3})
        self.assertEqual(list(frame["value"]), [2, 3, 4, 5, 6])
        self.assertEqual(list(frame["Campione"]), [1, 2, 3, 4, 5])

    def test_disabled_datasets_are_skipped(self):
        a = {"data": _df(["2026-01-01 10:00"], [1]), "enabled": True}
        b = {"data": _df(["2026-01-02 10:00"], [9]), "enabled": False}
        frame = export.build_export_frame([a, b], {"trim_start": 0, "trim_end": 0})
        self.assertEqual(list(frame["value"]), [1])

    def test_combines_and_sorts_multiple_datasets(self):
        a = {"data": _df(["2026-01-02 10:00"], [2]), "enabled": True}
        b = {"data": _df(["2026-01-01 10:00"], [1]), "enabled": True}
        frame = export.build_export_frame([a, b], {"trim_start": 0, "trim_end": 0})
        self.assertEqual(list(frame["value"]), [1, 2])
        self.assertEqual(list(frame["Campione"]), [1, 2])

    def test_trim_too_large_returns_empty(self):
        ds = {
            "data": _df(["2026-01-01 10:00", "2026-01-01 10:15"], [1, 2]),
            "enabled": True,
        }
        frame = export.build_export_frame([ds], {"trim_start": 1, "trim_end": 1})
        self.assertTrue(frame.empty)


class TestColumnSelection(unittest.TestCase):
    def _frame(self):
        ds = {
            "data": pd.DataFrame(
                {
                    "Time": pd.to_datetime(["2026-01-01 10:00", "2026-01-01 10:15"]),
                    "Turbidity NTU": [1.0, 2.0],
                    "DO2": [3.0, 4.0],
                }
            ),
            "enabled": True,
        }
        return export.build_export_frame([ds], {"trim_start": 0, "trim_end": 0})

    def test_fixed_columns_are_not_selectable(self):
        frame = self._frame()
        selectable = export.selectable_columns(frame)
        self.assertEqual(selectable, ["Turbidity NTU", "DO2"])
        self.assertNotIn("Campione", selectable)
        self.assertNotIn("Time", selectable)

    def test_subset_keeps_fixed_columns_and_selected_only(self):
        frame = self._frame()
        out = export.subset_export(frame, ["DO2"])
        self.assertEqual(list(out.columns), ["Campione", "Time", "DO2"])

    def test_subset_default_empty_exports_only_fixed_columns(self):
        frame = self._frame()
        out = export.subset_export(frame, [])
        self.assertEqual(list(out.columns), ["Campione", "Time"])

    def test_subset_ignores_unknown_and_fixed_names(self):
        frame = self._frame()
        out = export.subset_export(frame, ["DO2", "Campione", "Nope", "Time"])
        self.assertEqual(list(out.columns), ["Campione", "Time", "DO2"])

    def test_rows_without_selected_data_are_dropped(self):
        ds = {
            "data": pd.DataFrame(
                {
                    "Time": pd.to_datetime(
                        ["2026-01-01 10:00", "2026-01-01 10:01", "2026-01-01 10:02", "2026-01-01 10:03"]
                    ),
                    "Turbidity NTU": [1.0, None, 3.0, None],
                    "DO2": [None, 2.0, None, None],
                }
            ),
            "enabled": True,
        }
        frame = export.build_export_frame([ds], {"trim_start": 0, "trim_end": 0})

        out = export.subset_export(frame, ["Turbidity NTU"])
        self.assertEqual(list(out["Turbidity NTU"]), [1.0, 3.0])
        self.assertEqual(list(out["Campione"]), [1, 2])

        out_both = export.subset_export(frame, ["Turbidity NTU", "DO2"])
        # Row 10:03 has neither column -> dropped; the others survive.
        self.assertEqual(list(out_both["Campione"]), [1, 2, 3])
        self.assertEqual(out_both["Turbidity NTU"].dropna().tolist(), [1.0, 3.0])
        self.assertEqual(out_both["DO2"].dropna().tolist(), [2.0])
        self.assertTrue(pd.isna(out_both["DO2"].iloc[0]))
        self.assertTrue(pd.isna(out_both["DO2"].iloc[2]))

    def test_renumber_after_filtering_keeps_time_order(self):
        ds = {
            "data": pd.DataFrame(
                {
                    "Time": pd.to_datetime(
                        ["2026-01-01 10:00", "2026-01-01 10:01", "2026-01-01 10:02", "2026-01-01 10:03"]
                    ),
                    "DO2": [None, 1.0, None, 2.0],
                }
            ),
            "enabled": True,
        }
        frame = export.build_export_frame([ds], {"trim_start": 0, "trim_end": 0})
        out = export.subset_export(frame, ["DO2"])
        self.assertEqual(list(out["DO2"]), [1.0, 2.0])
        self.assertEqual(list(out["Campione"]), [1, 2])
        self.assertEqual(list(out["Time"]), list(frame["Time"].iloc[[1, 3]]))


class TestThresholdColumns(unittest.TestCase):
    def _frame(self):
        return pd.DataFrame(
            {
                "Time": pd.to_datetime(
                    ["2026-01-01 10:00", "2026-01-01 10:01", "2026-01-01 10:02", "2026-01-01 10:03"]
                ),
                "Turbidity NTU": [2.0, 5.0, 8.0, None],
            }
        )

    def test_exceedance_values(self):
        frame = self._frame()
        cols = export.threshold_columns(frame, 4.0, 7.0)
        self.assertEqual(list(cols.columns), ["Soglia 4 NTU", "Soglia 7 NTU"])
        self.assertEqual(cols["Soglia 4 NTU"].dropna().tolist(), [0.0, 1.0, 1.0])
        self.assertEqual(cols["Soglia 7 NTU"].dropna().tolist(), [0.0, 0.0, 1.0])
        self.assertTrue(pd.isna(cols["Soglia 4 NTU"].iloc[3]))
        self.assertTrue(pd.isna(cols["Soglia 7 NTU"].iloc[3]))

    def test_threshold_column_names_use_actual_values(self):
        frame = self._frame()
        cols = export.threshold_columns(frame, 4.5, 7.0)
        self.assertEqual(list(cols.columns), ["Soglia 4.5 NTU", "Soglia 7 NTU"])

    def test_threshold_is_inclusive(self):
        frame = pd.DataFrame(
            {
                "Time": pd.to_datetime(["2026-01-01 10:00", "2026-01-01 10:01"]),
                "Turbidity NTU": [4.0, 7.0],
            }
        )
        cols = export.threshold_columns(frame, 4.0, 7.0)
        self.assertEqual(cols["Soglia 4 NTU"].tolist(), [1.0, 1.0])
        self.assertEqual(cols["Soglia 7 NTU"].tolist(), [0.0, 1.0])

    def test_subset_export_appends_threshold_columns(self):
        ds = {
            "data": pd.DataFrame(
                {
                    "Time": pd.to_datetime(
                        ["2026-01-01 10:00", "2026-01-01 10:01", "2026-01-01 10:02"]
                    ),
                    "Turbidity NTU": [2.0, 5.0, 8.0],
                }
            ),
            "enabled": True,
        }
        frame = export.build_export_frame([ds], {"trim_start": 0, "trim_end": 0})
        out = export.subset_export(frame, ["Turbidity NTU"], thresholds=(4.0, 7.0))
        self.assertEqual(
            list(out.columns),
            ["Campione", "Time", "Turbidity NTU", "Soglia 4 NTU", "Soglia 7 NTU"],
        )
        self.assertEqual(out["Soglia 4 NTU"].tolist(), [0.0, 1.0, 1.0])
        self.assertEqual(out["Soglia 7 NTU"].tolist(), [0.0, 0.0, 1.0])

    def test_write_export_xlsx_highlights_threshold_rows(self):
        from io import BytesIO

        from openpyxl import load_workbook

        ds = {
            "data": pd.DataFrame(
                {
                    "Time": pd.to_datetime(
                        ["2026-01-01 10:00", "2026-01-01 10:01", "2026-01-01 10:02"]
                    ),
                    "Turbidity NTU": [3.0, 5.0, 8.0],
                }
            ),
            "enabled": True,
        }
        frame = export.build_export_frame([ds], {"trim_start": 0, "trim_end": 0})
        out = export.subset_export(frame, ["Turbidity NTU"], thresholds=(4.0, 7.0))

        buf = BytesIO()
        export.write_export_xlsx(out, (4.0, 7.0), buf)
        buf.seek(0)
        ws = load_workbook(buf)["Dati Raw"]

        def fill_of(row):
            return ws.cell(row=row, column=1).fill

        # 3.0 NTU -> no alert
        self.assertIsNone(fill_of(2).fill_type)
        # 5.0 NTU -> soglia 1 only -> yellow
        self.assertEqual(fill_of(3).fill_type, "solid")
        self.assertTrue(fill_of(3).start_color.rgb.endswith("FFFF00"))
        # 8.0 NTU -> soglia 2 -> red
        self.assertEqual(fill_of(4).fill_type, "solid")
        self.assertTrue(fill_of(4).start_color.rgb.endswith("FF0000"))


if __name__ == "__main__":
    unittest.main()
