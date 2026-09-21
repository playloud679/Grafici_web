"""Raw measurement export helpers.

Builds a single combined worksheet of raw measurement rows, after the
synchronous initial/final row trim configured in the UI, and prepends a
progressive sample number assigned in time order.

Module documentation: ``docs/modules/export.md``.
"""

from __future__ import annotations

import pandas as pd

# Column added to the export; values are 1..n from the first acquired sample.
SAMPLE_COLUMN = "Campione"

# Columns always exported: sample number and timestamp. Everything else is
# toggleable by the user (default disabled in the UI).
FIXED_COLUMNS = (SAMPLE_COLUMN, "Time")

# Turbidity column used to compute the guard-threshold exceedance columns.
TURBIDITY_COLUMN = "Turbidity NTU"


def _trimmed(df: pd.DataFrame, trim_start: int, trim_end: int) -> pd.DataFrame | None:
    """Apply the synchronous row trim; return ``None`` if it empties the report.

    Mirrors the trim used by the chart (``plotting._trim``) so the exported
    rows match exactly what is plotted. Kept here to avoid an import of the
    heavy matplotlib plotting module from the export path.
    """
    if trim_start + trim_end >= len(df):
        return None
    if trim_end > 0:
        return df.iloc[trim_start : len(df) - trim_end]
    return df.iloc[trim_start:]


def build_export_frame(datasets, trim_params) -> pd.DataFrame:
    """Combine enabled datasets into one raw export frame.

    Each enabled report is trimmed by ``trim_start`` / ``trim_end`` rows
    (the same synchronous trim used by every chart), then all remaining rows
    are concatenated and sorted by ``Time``. A ``Campione`` column is prepended
    with values ``1..n`` starting from the first acquired sample in time order.

    Returns an empty DataFrame when there is no enabled data left after the
    trim.
    """
    trim_start = max(0, int(trim_params.get("trim_start", 0)))
    trim_end = max(0, int(trim_params.get("trim_end", 0)))

    frames: list[pd.DataFrame] = []
    for ds in datasets:
        if not ds.get("enabled", True):
            continue
        trimmed = _trimmed(ds["data"], trim_start, trim_end)
        if trimmed is not None and not trimmed.empty:
            frames.append(trimmed)

    if not frames:
        return pd.DataFrame()

    combined = pd.concat(frames, ignore_index=True)
    if "Time" in combined.columns:
        combined = combined.sort_values("Time").reset_index(drop=True)

    combined.insert(0, SAMPLE_COLUMN, range(1, len(combined) + 1))
    return combined


def selectable_columns(frame: pd.DataFrame) -> list[str]:
    """Return the data columns the user can toggle on/off.

    ``Campione`` and ``Time`` are fixed and therefore excluded from the list.
    """
    return [c for c in frame.columns if c not in FIXED_COLUMNS]


def _format_threshold(value: float) -> str:
    """Render a threshold without a trailing ``.0`` (``4.0`` -> ``"4"``)."""
    if float(value).is_integer():
        return str(int(value))
    return f"{value:g}"


def threshold_column_names(low: float, high: float) -> tuple[str, str]:
    """Return the two exceedance column names for the given thresholds."""
    return (
        f"Soglia {_format_threshold(low)} NTU",
        f"Soglia {_format_threshold(high)} NTU",
    )


def threshold_columns(frame: pd.DataFrame, low: float, high: float) -> pd.DataFrame:
    """Return two guard-threshold exceedance columns for turbidity.

    ``Soglia {low} NTU`` is 1 where turbidity is ``>=`` ``low`` and 0 where it
    is not; ``Soglia {high} NTU`` is analogous for ``high``. Missing turbidity
    values yield NaN (empty cells) in both columns.
    """
    if TURBIDITY_COLUMN in frame.columns:
        turb = pd.to_numeric(frame[TURBIDITY_COLUMN], errors="coerce")
    else:
        turb = pd.Series(index=frame.index, dtype=float)

    names = threshold_column_names(low, high)
    result = pd.DataFrame(index=frame.index)
    for name, threshold in zip(names, (low, high)):
        result[name] = (turb >= threshold).astype(float).where(turb.notna())
    return result


def subset_export(frame: pd.DataFrame, selected_columns, thresholds=None) -> pd.DataFrame:
    """Return ``Campione`` + ``Time`` + the chosen data columns, in order.

    Unknown or fixed column names in ``selected_columns`` are ignored, so the
    fixed sample/time columns can never be duplicated or dropped. When at least
    one data column is selected, rows where **every** selected column is empty
    (all NaN) are dropped — the interleaved AquaScan format otherwise leaves
    most rows blank for any single sensor — and ``Campione`` is renumbered
    ``1..n`` in time order over the surviving rows.

    ``thresholds`` optionally provides ``(low, high)`` NTU guard thresholds;
    two exceedance columns are then appended (see :func:`threshold_columns`).
    """
    selected = [c for c in selected_columns if c in frame.columns and c not in FIXED_COLUMNS]
    out = frame[list(FIXED_COLUMNS) + selected]

    if thresholds is not None:
        low, high = thresholds
        out = pd.concat([out, threshold_columns(frame, low, high)], axis=1)

    if selected:
        out = out.dropna(subset=selected, how="all")
    out = out.reset_index(drop=True)
    out[SAMPLE_COLUMN] = range(1, len(out) + 1)
    return out


def write_export_xlsx(frame: pd.DataFrame, thresholds, buffer) -> None:
    """Write ``frame`` to ``buffer`` as XLSX with threshold row highlighting.

    Rows where the second threshold column is 1 are filled red; rows where only
    the first threshold column is 1 are filled yellow. The fill spans the whole
    row so an alert is immediately visible across all exported columns.
    """
    from openpyxl.styles import PatternFill

    low, high = thresholds
    low_col, high_col = threshold_column_names(low, high)
    yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    red = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Dati Raw")
        ws = writer.sheets["Dati Raw"]

        header_to_col = {cell.value: cell.column for cell in ws[1]}
        low_idx = header_to_col.get(low_col)
        high_idx = header_to_col.get(high_col)
        if low_idx is None and high_idx is None:
            return

        for row in range(2, ws.max_row + 1):
            fill = None
            if high_idx is not None and ws.cell(row=row, column=high_idx).value == 1:
                fill = red
            elif low_idx is not None and ws.cell(row=row, column=low_idx).value == 1:
                fill = yellow
            if fill is not None:
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = fill
